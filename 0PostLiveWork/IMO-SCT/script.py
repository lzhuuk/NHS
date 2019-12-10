import openpyxl


workbook1 = openpyxl.load_workbook('ICDWHO_LEXICALS_TEXT_IMO.xlsx')
worksheet1 = workbook1.worksheets[0]
workbook2 = openpyxl.load_workbook('output.xlsx')
worksheet2 = workbook2.worksheets[0]
workbook3 = openpyxl.load_workbook('database.xlsx')
worksheet3 = workbook3.worksheets[0]
workbook4 = openpyxl.load_workbook('ICDWHO_IMO.xlsx')
worksheet4 = workbook4.worksheets[0]
workbook5 = openpyxl.load_workbook('ICDWHO_BASE_TEXT.xlsx')
worksheet5 = workbook5.worksheets[0]
max_row1 = worksheet1.max_row
max_row2 = worksheet2.max_row
max_row3 = worksheet3.max_row
max_row4 = worksheet4.max_row
max_row5 = worksheet5.max_row
IMO_ID = []
SCT_ID = []
ICD_ID = []
ICD_ID1 = []
for x in range(2, max_row2 + 1):
    IMO_ID.append(str(worksheet2.cell(row=x, column=1).value))
    SCT_ID.append(str(worksheet2.cell(row=x, column=3).value))
IMO_ID1 = []
SCT_ID1 = []
for y in range(2, max_row1 + 1):
    IMO_ID1.append(str(worksheet1.cell(row=y, column=1).value))
for z in range(2, max_row3 + 1):
    SCT_ID1.append(worksheet3.cell(row=z, column=1).value)
for a in range(2, max_row4 + 1):
    ICD_ID.append(str(worksheet4.cell(row=a, column=1).value))
for b in range(2, max_row5 + 1):
    ICD_ID1.append(worksheet5.cell(row=b, column=1).value)
worksheet2.cell(row=1, column=2).value = 'IMO_TEXT'
worksheet2.cell(row=1, column=4).value = 'SCT_CONCEPT_TERM'
worksheet2.cell(row=1, column=5).value = 'ICD_CODE'
worksheet2.cell(row=1, column=6).value = 'ICD_TEXT'
for i in range(0, len(IMO_ID)):
    if i % 1000 == 0:
        print(i)
    if IMO_ID[i] in IMO_ID1:
        # print("In")
        index = IMO_ID1.index(IMO_ID[i])
        name = worksheet1.cell(row=index + 2, column=2).value
        worksheet2.cell(row=i + 2, column=2).value = name
    elif IMO_ID[i] not in IMO_ID1:
        # print("Out")
        worksheet2.cell(row=i + 2, column=2).value = ''
    elif IMO_ID[i] in ICD_ID:
        index0 = ICD_ID.index(IMO_ID[i])
        code = worksheet4.cell(row=index0 + 2, column=3).value
        if code in ICD_ID1:
            index2 = ICD_ID1.index(code)
            name0 = worksheet5.cell(row=index2 + 2, column=2).value
            worksheet2.cell(row=i + 2, column=6).value = name0
        else:
            worksheet2.cell(row=i + 2, column=6).value = ''
        # name0 = worksheet4.cell(row=index0 + 2, column=3).value
        worksheet2.cell(row=i + 2, column=5).value = code
        # worksheet2.cell(row=i + 2, column=6).value = name0
    else:
        worksheet2.cell(row=i + 2, column=5).value = ''
        # worksheet2.cell(row=i + 2, column=6).value = ''

for j in range(0, len(SCT_ID)):
    if j % 1000 == 0:
        print(j)
    if SCT_ID[j] in SCT_ID1:
        # print("In")
        index1 = SCT_ID1.index(SCT_ID[j])
        term = worksheet3.cell(row=index1 + 2, column=2).value
        worksheet2.cell(row=j + 2, column=4).value = term
    else:
        # print("Out")
        worksheet2.cell(row=j + 2, column=4).value = ''

workbook2.save('output.xlsx')
